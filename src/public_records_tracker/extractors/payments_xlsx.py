from __future__ import annotations

from io import BytesIO
from urllib.parse import urlsplit

import duckdb
from openpyxl import load_workbook

from ..models import Record
from ..structured import add_fact, upsert_entity
from ..supplier_hygiene import split_payment_channel
from .payments import (
    _amount,
    _date,
    _field,
    _is_supplier_file,
    _normalise_header,
    _normalised_aliases,
    _payer_entity,
)


def _is_xlsx(record: Record) -> bool:
    content_type = record.content_type.casefold()
    path = urlsplit(record.url).path.casefold()
    metadata_name = " ".join(
        [record.title, str(record.metadata.get("anchor_text") or "")]
    ).casefold()
    return (
        path.endswith((".xlsx", ".xlsm"))
        or ".xlsx" in metadata_name
        or ".xlsm" in metadata_name
        or "spreadsheetml" in content_type
    )


def _sheet_rows(sheet) -> tuple[int, list[str], list[tuple[int, dict[str, str]]]] | None:
    supplier_aliases = _normalised_aliases("supplier")
    amount_aliases = _normalised_aliases("amount")
    rows = list(sheet.iter_rows(values_only=True))
    for offset, values in enumerate(rows[:25], start=1):
        headers = [str(value or "").strip() for value in values]
        normalized = {_normalise_header(value) for value in headers if value}
        if not (normalized & supplier_aliases and normalized & amount_aliases):
            continue
        data: list[tuple[int, dict[str, str]]] = []
        for row_number, row_values in enumerate(rows[offset:], start=offset + 1):
            row = {
                header: str(value).strip() if value is not None else ""
                for header, value in zip(headers, row_values, strict=False)
                if header
            }
            if any(row.values()):
                data.append((row_number, row))
        return offset, headers, data
    return None


def extract_payments_xlsx(
    *,
    con: duckdb.DuckDBPyConnection,
    record: Record,
    document_id: str,
    snapshot_id: str,
    source: dict,
) -> int:
    if not _is_xlsx(record):
        return 0

    try:
        workbook = load_workbook(BytesIO(record.body), read_only=True, data_only=True)
    except Exception:
        return 0

    payer = _payer_entity(con, source, record)
    count = 0
    try:
        for sheet in workbook.worksheets:
            parsed = _sheet_rows(sheet)
            if parsed is None:
                continue
            _, _, rows = parsed
            # Header detection is deliberately required before the source-level
            # allow-by-header rule can admit a generic council download.
            if not _is_supplier_file(record, source, has_headers=True):
                continue
            for row_number, row in rows:
                raw_supplier_name = _field(row, "supplier")
                supplier_name, payment_channel = split_payment_channel(raw_supplier_name)
                amount = _amount(_field(row, "amount"))
                if not supplier_name or amount is None:
                    continue
                supplier = upsert_entity(
                    con,
                    entity_type="SUPPLIER",
                    name=supplier_name,
                    namespace=f"payments:{source['id']}",
                    metadata={"source_system": source.get("name", source["id"])},
                )
                payment_date = _date(_field(row, "date"))
                description = _field(row, "description") or None
                department = _field(row, "department") or None
                reference = _field(row, "reference") or None
                locator = f"XLSX {sheet.title}!row {row_number}"
                add_fact(
                    con,
                    document_id=document_id,
                    snapshot_id=snapshot_id,
                    fact_type="PAYMENT",
                    predicate="PAYMENT_TO_SUPPLIER",
                    evidence_class=record.evidence_class.value,
                    subject_entity_id=payer,
                    object_entity_id=supplier,
                    value_text=amount,
                    locator=locator,
                    metadata={
                        "currency": "GBP",
                        "payment_date": payment_date,
                        "description": description,
                        "department": department,
                        "reference": reference,
                        "source_row": row_number,
                        "sheet": sheet.title,
                        "raw_supplier_name": raw_supplier_name or supplier_name,
                        "payment_channel": payment_channel,
                    },
                )
                count += 1
    finally:
        workbook.close()
    return count
